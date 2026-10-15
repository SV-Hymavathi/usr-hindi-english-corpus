#!/usr/bin/env python3
"""Meaning-based metrics for every complete system: COMET (Unbabel/wmt22-comet-da, reference-based, uses the Hindi source)
and BERTScore F1 (roberta-large, English, baseline-rescaled). Resumable: per-system scores are cached in semantic_cache/.

Reproduce (Linux / WSL, CPU is enough; first run downloads the models: COMET ~2.3 GB, roberta-large ~1.4 GB):
    python3 -m venv ~/semenv && ~/semenv/bin/pip install "unbabel-comet>=2.2" bert-score openpyxl sacrebleu
    ~/semenv/bin/python semantic_metrics.py --metric bertscore
    ~/semenv/bin/python semantic_metrics.py --metric comet
    ~/semenv/bin/python semantic_metrics.py --report          # -> semantic_metrics_results.xlsx / .tsv
Options: --batch 16 (lower it if memory is short), --only exp1:google_translate (one system)."""
import argparse, json, os, sys
HERE = os.path.dirname(os.path.abspath(__file__))
EXP = os.path.join(HERE, '..')
CACHE = os.path.join(HERE, 'semantic_cache'); os.makedirs(CACHE, exist_ok=True)
ap = argparse.ArgumentParser()
ap.add_argument('--metric', choices=['comet', 'bertscore'])
ap.add_argument('--report', action='store_true')
ap.add_argument('--batch', type=int, default=16)
ap.add_argument('--only', nargs='*')
a = ap.parse_args()
DOMS = ['tourism', 'health', 'geography']
ITEMS = []
for d in DOMS:
    for l in open(os.path.join(EXP, 'inputs', d + '.jsonl'), encoding='utf-8'):
        r = json.loads(l)
        if r['reference']:
            ITEMS.append((d, r['id'], r['hindi'], r['reference']))
SRC = {}
f1 = os.path.join(EXP, 'exp1', 'exp1_input.tsv')
for l in open(f1, encoding='utf-8'):
    p = l.rstrip('\n').split('\t')
    if len(p) == 3:
        SRC[(p[0], p[1])] = p[2]


def systems():
    out = {}
    for exp in ('exp1', 'exp2', 'exp3'):
        od = os.path.join(EXP, exp, 'outputs')
        if not os.path.isdir(od):
            continue
        for fn in sorted(os.listdir(od)):
            if not fn.endswith('.tsv'):
                continue
            o = {}
            for l in open(os.path.join(od, fn), encoding='utf-8'):
                p = l.rstrip('\n').split('\t', 2)
                if len(p) == 3:
                    o[(p[0], p[1])] = p[2]
            if all((d, i) in o for d, i, _, _ in ITEMS):
                out['%s:%s' % (exp, fn[:-4])] = o
    return out


SYS = systems()
if a.only:
    SYS = {k: v for k, v in SYS.items() if k in a.only}
print('complete systems:', list(SYS))

if a.metric == 'bertscore':
    from bert_score import score
    for s, o in SYS.items():
        cf = os.path.join(CACHE, 'bertscore__%s.json' % s.replace(':', '__'))
        if os.path.exists(cf):
            continue
        cands = [o[(d, i)] for d, i, _, _ in ITEMS]; refs = [r for _, _, _, r in ITEMS]
        P, R, F = score(cands, refs, lang='en', rescale_with_baseline=True, batch_size=a.batch, verbose=True)
        json.dump([float(x) for x in F], open(cf, 'w'))
        print(s, 'BERTScore F1 = %.4f' % float(F.mean()))
elif a.metric == 'comet':
    from comet import download_model, load_from_checkpoint
    model = load_from_checkpoint(download_model('Unbabel/wmt22-comet-da'))
    for s, o in SYS.items():
        cf = os.path.join(CACHE, 'comet__%s.json' % s.replace(':', '__'))
        if os.path.exists(cf):
            continue
        data = [dict(src=SRC.get((d, i), h), mt=o[(d, i)], ref=r) for d, i, h, r in ITEMS]
        out = model.predict(data, batch_size=a.batch, gpus=0)
        json.dump([float(x) for x in out.scores], open(cf, 'w'))
        print(s, 'COMET = %.4f' % out.system_score)

if a.report:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    rows = []
    for s in SYS:
        row = [s]
        for m in ('comet', 'bertscore'):
            cf = os.path.join(CACHE, '%s__%s.json' % (m, s.replace(':', '__')))
            sc = json.load(open(cf)) if os.path.exists(cf) else None
            for scope in ['all'] + DOMS:
                if sc is None:
                    row.append(''); continue
                v = [x for x, it in zip(sc, ITEMS) if scope == 'all' or it[0] == scope]
                row.append(round(sum(v) / len(v), 4))
        rows.append(row)
    hdr = ['System'] + ['%s %s' % (m, sc) for m in ('COMET', 'BERTScore-F1') for sc in ['all'] + DOMS]
    with open(os.path.join(HERE, 'semantic_metrics_results.tsv'), 'w', encoding='utf-8') as f:
        f.write('\t'.join(hdr) + '\n')
        for r in rows:
            f.write('\t'.join(map(str, r)) + '\n')
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='305496')
    for r in rows:
        ws.append(r)
    ws.column_dimensions['A'].width = 40
    wb.save(os.path.join(HERE, 'semantic_metrics_results.xlsx'))
    for r in rows:
        print(r)
