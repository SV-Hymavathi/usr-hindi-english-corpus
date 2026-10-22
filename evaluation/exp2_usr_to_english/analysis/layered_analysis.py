#!/usr/bin/env python3
"""Layer-by-layer error analysis of Experiment 1 (Hindi -> English), guided by the USR. No LLM is used.

  python3 layered_analysis.py [--domain tourism]   ->  <domain>_error_analysis.xlsx  (one tab per layer)

Principle: the USR states what the sentence means at each layer; the gold reference confirms how English expresses it.
An output is flagged only when the USR and the reference AGREE and the output does not. Every flag is a candidate for
human confirmation (columns 'Human verdict' / 'Human note' on every tab).

Tabs / layers
 L1a Lexical & NE        lexical choice differs from the reference; named entity dropped / translated / replaced; number
 L1b Cultural concepts   culture-specific concept (resources/cultural_concepts.tsv) rendered literally or generically
                         (e.g. snAna 'holy dip' -> 'main bath', prakRAla 'ablution' -> 'foot washing', xarSana -> 'vision')
 L1c Synonyms            output used another dictionary sense of the same Hindi word that is a synonym of the reference
                         word (both listed together as senses of other Hindi words too): acceptable variant, not an error
 L1d Other sense used    output used another dictionary sense of the same Hindi word that is not listed as a synonym:
                         either polysemy (wrong sense) or a near-synonym; decide in the 'Human verdict' column
 L2  Morphology          plural / degree (superlative, comparative) / tense / aspect (perfect, progressive) / modality
 L3  Semantic relations  voice (passive <-> active) / negation lost or added / cause, purpose, without, like not expressed
 L4  Discourse           discourse relation of the USR (and, but, because, therefore, if-then, when, that ...) not expressed
 L5  Speaker's view      particles hI (only/very), BI (also), lagaBaga (about) ... lost; deixis this/that/here/there lost or
                         reversed; honorific (respect: Shri / ji / Lord / Mother) dropped
 L6  Sentence type       question / command / exclamation changed
 L7  Adequacy            omission (low coverage of reference content words) / addition (content not in reference)
 L8  Output form         hallucinated numbers, notes / markdown, Hindi script left, repetition, much longer / shorter
"""
import argparse, collections, difflib, itertools, json, os, re, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
import sacrebleu

HERE = os.path.dirname(os.path.abspath(__file__))
RES = os.path.join(HERE, 'resources')
ap = argparse.ArgumentParser(); ap.add_argument('--domain', default='tourism'); a = ap.parse_args()
D = a.domain
CFG = json.load(open(os.path.join(HERE, 'config.json'), encoding='utf-8'))
items = [json.loads(l) for l in open(os.path.join(HERE, '..', 'inputs', D + '.jsonl'), encoding='utf-8')]
hindi = {l.split('\t')[1]: l.rstrip('\n').split('\t')[2] for l in open(os.path.join(HERE, 'exp1_input.tsv'), encoding='utf-8')
         if l.startswith(D + '\t')}
SYS = {}
for n in CFG['systems']:
    f = os.path.join(HERE, 'outputs', n + '.tsv')
    if os.path.exists(f):
        o = {l.split('\t')[1]: l.rstrip('\n').split('\t', 2)[2] for l in open(f, encoding='utf-8') if l.startswith(D + '\t') and l.count('\t') >= 2}
        if len(o) >= 0.95 * len(items):
            SYS[n] = o
if not SYS:
    sys.exit('no system has finished the %s domain yet' % D)

# ---------------- resources ----------------
EN, SENSES = {}, collections.defaultdict(set)
for l in open(os.path.join(RES, 'concept-to-mrs-rels.dat'), encoding='utf-8'):
    c = l.split()
    if len(c) >= 4 and re.search(r'_\d+$', c[0]) and '+' not in c[0]:
        e = re.sub(r'_\d+', '', c[3]).replace('+', ' ').lower()
        EN[c[0]] = e
        if e and e != '-' and len(e) > 2:
            SENSES[re.sub(r'_\d+$', '', c[0])].add(e)
PAIR = collections.Counter()                     # how many Hindi roots list both English words as senses
for root, es in SENSES.items():
    if 2 <= len(es) <= 40:
        for x, y in itertools.combinations(sorted(es), 2):
            PAIR[(x, y)] += 1


def synonyms(x, y):
    return PAIR[tuple(sorted((x, y)))] >= 2


TAM = {}
for l in open(os.path.join(RES, 'tam.tsv'), encoding='utf-8'):
    c = l.rstrip('\n').split('\t')
    if len(c) >= 4 and not l.startswith('#'):
        TAM[c[0]] = (c[1], c[2], c[3])
CULT = {}
for l in open(os.path.join(RES, 'cultural_concepts.tsv'), encoding='utf-8'):
    if l.startswith('#') or not l.strip():
        continue
    c = l.rstrip('\n').split('\t') + ['', '', '']
    CULT[c[0]] = (c[1], c[2], c[3])

STOP = set('a an the of in on at to by with from for and or but is are was were be been being am has have had do does did this '
           'that these those it its as than so very also only he she they we you i his her their our your my me him them us which '
           'who whom whose what when where how why there here will would can could may might shall should must into up out about '
           'over after before then just such any all some each other more most own same too'.split())


def stem(w):
    for suf in ('ing', 'ies', 'ed', 'es', 's', 'ly'):
        if len(w) > len(suf) + 2 and w.endswith(suf):
            return w[:-len(suf)]
    return w


def toks(s):
    return re.findall(r"[a-z0-9']+", s.lower())


def has_any(s, pats):
    return re.search(r'\b(' + '|'.join(pats) + r')\b', s.lower()) is not None


def has_word(text, en):
    ws = [stem(w) for w in toks(en) if w not in STOP]
    tset = {stem(w) for w in toks(text)}
    return bool(ws) and all(w in tset for w in ws)


WX = [('A', 'a'), ('I', 'i'), ('U', 'u'), ('q', 'ri'), ('E', 'ai'), ('O', 'au'), ('M', 'n'), ('z', 'n'), ('H', ''), ('K', 'kh'),
      ('G', 'gh'), ('f', 'n'), ('C', 'chh'), ('c', 'ch'), ('J', 'jh'), ('F', 'n'), ('T', 'th'), ('D', 'dh'), ('N', 'n'),
      ('W', 'th'), ('w', 't'), ('X', 'dh'), ('x', 'd'), ('P', 'ph'), ('B', 'bh'), ('S', 'sh'), ('R', 'sh'), ('Y', '')]


def wx2lat(s):
    return ''.join(dict(WX).get(ch, ch) for ch in re.sub(r'_\d+$', '', s)).lower()


def skel(s):
    s = s.lower().replace('sh', 's').replace('ph', 'f').replace('w', 'v').replace('ee', 'i').replace('oo', 'u')
    return re.sub(r'(.)\1+', r'\1', re.sub(r'[aeiouyh\W\d_]', '', s))


def present(name, text):
    sk = skel(name)
    if len(sk) < 2:
        return True
    tt = re.findall(r'[A-Za-z]+', text)
    return any(difflib.SequenceMatcher(None, sk, skel(''.join(tt[i:i + n]))).ratio() >= 0.75
               for n in (1, 2, 3, 4) for i in range(len(tt) - n + 1))


def parse_usr(usr):
    nodes, stype = [], ''
    for l in usr.splitlines():
        l = l.strip()
        if l.startswith('%'):
            stype = l; continue
        c = l.split()
        if len(c) >= 9 and re.fullmatch(r'\d+', c[1]):
            concept, tam = c[0], ''
            m = re.match(r'^(.*?_\d+)-(.+)$', concept)
            if m and not concept.startswith('['):
                concept, tam = m.group(1), m.group(2)
            nodes.append(dict(concept=concept, root=re.sub(r'_\d+$', '', concept), idx=c[1], sem=c[2], gnp=c[3],
                              rel=c[4].split(':')[-1] if ':' in c[4] else '', disc=c[5].split(':')[-1] if ':' in c[5] else '',
                              spk=c[6], scope=c[7], cx=c[8], tam=tam))
    return nodes, stype


PASSIVE = re.compile(r"\b(is|are|was|were|be|been|being|get|gets|got|gotten)\s+(\w+ly\s+)?(\w+ed|\w+en|built|made|said|known|held|"
                     r"kept|seen|found|told|given|taken|done|sold|brought|bought|thought|taught|put|set|left|led|laid|paid|"
                     r"spread|shown|called|considered|situated|located)\b", re.I)
STATIVE = r'\b(filled|situated|located|known|surrounded|spread|covered|dedicated|called|considered|believed)\b'
NEG_EN = re.compile(r"\b(not|no|never|none|nor|without|cannot|neither|nothing|nobody)\b|n't", re.I)
IRR_PAST = set('came went was were had did made took gave got told said saw became began brought bought thought taught stood led ran '
               'sat lay laid paid spent sent fell felt grew knew rose wrote drew'.split())
PRES_W = {'is', 'are', 'am', 'has', 'have', 'does', 'do', 'can', 'may', 'must', 'should'}
PAST_W = {'was', 'were', 'had', 'did', 'could', 'might', 'would'} | IRR_PAST


def tense(s):
    w = re.findall(r"[a-z']+", s.lower())
    if {'will', 'shall', "'ll"} & set(w):
        return 'future'
    pr, pa = PRES_W & set(w), PAST_W & set(w)
    if pr and not pa:
        return 'present'
    if pa and not pr:
        return 'past'
    if pr and pa:
        return 'present' if next(x for x in w if x in pr or x in pa) in pr else 'past'
    return ''


def perfect(s):
    return re.search(r'\b(has|have|had)\s+(\w+ly\s+)?(been\s+)?\w+(ed|en|n|t)\b', s.lower()) is not None


def progressive(s):
    return re.search(r'\b(is|are|was|were|am|be|been)\s+\w+ing\b', s.lower()) is not None


MODAL = {'ability': ['can', 'could', 'able', 'cannot'], 'obligation': ['should', 'must', 'have to', 'has to', 'had to', 'ought', 'need to', 'needs to'],
         'possibility': ['may', 'might', 'could', 'perhaps', 'probably']}
DISC = {'samuccaya': ['and', 'also', 'moreover', 'furthermore', 'besides', 'as well', 'too', 'in addition'],
        'viroXI': ['but', 'however', 'yet', 'though', 'although', 'whereas', 'while', 'still', 'nevertheless'],
        'viroXi': ['but', 'however', 'yet', 'though', 'although', 'whereas', 'while', 'still'],
        'viroXI_xyowaka': ['whereas', 'while', 'but', 'however'],
        'kAryakAraNa': ['because', 'since', 'as', 'due to', 'owing to', 'so', 'therefore', 'hence', 'thus'],
        'pariNAma': ['therefore', 'so', 'thus', 'hence', 'as a result', 'consequently', 'that is why', 'this is why'],
        'AvaSyakawApariNAma': ['if', 'then', 'so', 'only'],
        'uwwarakAla': ['then', 'after', 'later', 'afterwards', 'thereafter', 'subsequently'],
        'vyaBicAra': ['although', 'though', 'even', 'despite', 'yet', 'but', 'still', 'however'],
        'anyawra': ['or', 'otherwise', 'either', 'else'],
        'samAnakAla': ['when', 'while', 'as', 'whenever', 'during'],
        'uxxeSya': ['so that', 'in order', 'to', 'for'],
        'pUrvakAla': ['after', 'before', 'then', 'having'],
        'arWAw': ['that is', 'i.e', 'namely', 'means', 'meaning', 'in other words'],
        'vAkyakarma': ['that', 'whether', 'what', 'how', 'which', 'if', 'who', 'where', 'when', 'why']}
REL = {'rh': ['because', 'due to', 'owing to', 'as a result', 'since', 'for', 'on account'], 'rt': ['for', 'to', 'in order', 'so that'],
       'rasneg': ['without', 'not', 'no'], 'rsm': ['like', 'as', 'similar', 'such as']}
PART = {'hI': ['only', 'just', 'even', 'itself', 'very', 'alone', 'exactly', 'same', 'right', 'solely'],
        'BI': ['also', 'too', 'even', 'as well', 'both', 'either'],
        'lagaBaga': ['about', 'almost', 'nearly', 'around', 'approximately', 'roughly'],
        'kevala': ['only', 'just', 'merely', 'solely'], 'sirPa': ['only', 'just', 'merely'], 'mAwra': ['only', 'just', 'merely']}
DEIX = {'proximal': ['this', 'these', 'here'], 'distal': ['that', 'those', 'there']}
HONOR = r'\b(shri|sri|shree|ji|lord|mother|maa|mata|bhagwan|saint|sant|guru|swami|devi)\b'
DIG = str.maketrans('०१२३४५६७८९', '0123456789')

TABS = ['L1a Lexical & NE', 'L1b Cultural concepts', 'L1c Synonyms', 'L1d Other sense used', 'L2 Morphology', 'L3 Semantic relations',
        'L4 Discourse', "L5 Speaker's view", 'L6 Sentence type', 'L7 Adequacy', 'L8 Output form']
ERROR_TABS = [t for t in TABS if t != 'L1c Synonyms']          # synonyms are acceptable variants, not errors


def analyse(it, hin, ref, out):
    """returns list of (tab, sub-type, detail)"""
    R = []
    add = lambda t, s, d: R.append((t, s, d))
    nodes, st = parse_usr(it['usr'])
    if not out.strip():
        add('L8 Output form', 'empty output', ''); return R
    rl, ol = ref.lower(), out.lower()
    rset, oset = {stem(w) for w in toks(ref)}, {stem(w) for w in toks(out)}
    ne_heads = {n['idx'] for n in nodes if n['concept'].startswith('[ne')}
    names = collections.OrderedDict()
    for n in nodes:
        m = re.match(r'^(\d+):(begin|inside)$', n['cx'])
        if m and m.group(1) in ne_heads:
            names.setdefault(m.group(1), []).append(wx2lat(n['concept']))
            continue
        c = n['concept']
        if c.startswith(('[', '$', '@')) or not re.search(r'_\d+$', c) or n['tam']:
            continue
        # ---- cultural concepts
        if n['root'] in CULT:
            good, bad, note = CULT[n['root']]
            wrong = re.search(bad, ol) if bad else None
            loose = bad.replace(r'\b', '') if bad else ''
            explicit = wrong is not None and re.search(r'\(\s*[^)]*' + re.escape(wrong.group(0)), ol) is not None   # '(foot washing)'
            if wrong and not re.search(loose, rl) and (explicit or not re.search(good, ol)):
                add('L1b Cultural concepts', 'culturally wrong gloss', '%s: output "%s" (%s)' % (n['root'], wrong.group(0), note))
                continue
            if good and re.search(good, rl) and not re.search(good, ol):
                lit = re.search(bad, ol) if bad else None
                add('L1b Cultural concepts', 'literal / generic rendering' if lit else 'cultural concept missing',
                    '%s: reference "%s", output %s  (%s)' % (n['root'], re.search(good, rl).group(0),
                                                             ('"' + lit.group(0) + '"') if lit else 'has no equivalent', note))
                continue
        en = EN.get(c, '')
        if not en or en == '-' or len(en) < 3:
            continue
        if has_word(ref, en) and not has_word(out, en):
            others = [e for e in sorted(SENSES.get(n['root'], ())) if e != en and has_word(out, e) and not has_word(ref, e)]
            if others:
                o = others[0]
                if o[:4] == en[:4]:
                    continue                                     # word-form variant (hill / hilly), not a lexical issue
                if synonyms(en, o):
                    add('L1c Synonyms', 'synonym (acceptable variant)', '%s: reference "%s", output "%s"' % (n['root'], en, o))
                else:
                    add('L1d Other sense used', 'other sense of the Hindi word (polysemy or near-synonym: check)', '%s: intended "%s", output "%s"' % (n['root'], en, o))
            else:
                add('L1a Lexical & NE', 'lexical choice differs', '"%s" (%s) in reference, not in output' % (en, c))
    for nm in (' '.join(v) for v in names.values()):
        words = [w for w in nm.split() if len(skel(w)) >= 2] or [nm]
        if any(present(w, out) for w in words[:2]):
            continue
        if not any(present(w, ref) for w in words):
            caps = [w for k, w in enumerate(re.findall(r"[A-Za-z][A-Za-z'-]*", ref)) if k > 0 and w[0].isupper() and w.lower() not in STOP]
            if not caps or any(w.lower() in ol for w in caps):
                continue
        ocaps = [w for k, w in enumerate(re.findall(r"[A-Za-z][A-Za-z'-]*", out)) if k > 0 and w[0].isupper() and w.lower() not in STOP
                 and w.lower() not in rl and not present(w, ref)]
        if ocaps:
            add('L1a Lexical & NE', 'NE replaced', '%s -> %s' % (nm, ocaps[0]))
        else:
            add('L1a Lexical & NE', 'NE dropped / translated', nm)
    hn = sorted(set(re.findall(r'\d+(?:\.\d+)?', hin.translate(DIG).replace(',', ''))))
    on = sorted(set(re.findall(r'\d+(?:\.\d+)?', out.replace(',', ''))))
    if hn and [x for x in hn if x not in on]:
        add('L1a Lexical & NE', 'number missing / changed', 'Hindi %s, output %s' % (hn, on))
    extra = [x for x in on if x not in hn and x not in re.findall(r'\d+', ref)]
    if extra:
        add('L8 Output form', 'number added (hallucination)', str(extra))
    # ---- L2 morphology
    for n in nodes:
        en = EN.get(n['concept'], '')
        if not en or ' ' in en or len(en) < 3 or n['concept'].startswith(('[', '$')):
            continue
        pls = {en + 's', en + 'es', en[:-1] + 'ies'}
        if 'pl' in n['gnp'] and any(p in rl.split() for p in pls) and re.search(r'\b%s\b' % re.escape(en), ol) \
                and not any(re.search(r'\b%s\b' % re.escape(p), ol) for p in pls):
            add('L2 Morphology', 'number (plural lost)', en)
        if 'supl' in n['gnp'] and re.search(r'\w+est\b|\bmost\b|\bbest\b|\bworst\b|\bleast\b', rl) and \
                not re.search(r'\w+est\b|\bmost\b|\bbest\b|\bworst\b|\bleast\b', ol):
            add('L2 Morphology', 'degree (superlative lost)', en)
        if 'compr' in n['gnp'] and re.search(r'\w+er than|\bmore\b|\bless\b|\bbetter\b', rl) and \
                not re.search(r'\w+er\b|\bmore\b|\bless\b|\bbetter\b|\bworse\b', ol):
            add('L2 Morphology', 'degree (comparative lost)', en)
    tams = [TAM.get(n['tam']) for n in nodes if n['tam'] and n['rel'] == 'main']
    tams = [t for t in tams if t]
    if tams:
        t_te, t_as, t_mo = tams[0]
        if t_te in ('past', 'present', 'future') and tense(ref) == t_te and tense(out) not in (t_te, ''):
            add('L2 Morphology', 'tense', 'USR/reference %s, output %s' % (t_te, tense(out)))
        if t_as == 'perfect' and perfect(ref) and not perfect(out) and tense(out) != 'past':
            add('L2 Morphology', 'aspect (perfect lost)', '')
        if t_as in ('progressive', 'continuous') and progressive(ref) and not progressive(out):
            add('L2 Morphology', 'aspect (progressive lost)', '')
        mod = 'ability' if t_mo in ('can', 'could') else 'obligation' if t_mo in ('should', 'must', 'have to', 'has to') else \
            'possibility' if t_mo in ('may', 'might') else ''
        if mod and has_any(ref, MODAL[mod]) and not has_any(out, MODAL[mod] + ['able']):
            add('L2 Morphology', 'modality (%s lost)' % mod, t_mo)
    # ---- L3 semantic relations
    usr_pass = 'pass' in st
    if usr_pass and PASSIVE.search(ref) and not PASSIVE.search(out):
        add('L3 Semantic relations', 'voice (passive -> active)', '')
    if not usr_pass and st and not PASSIVE.search(ref) and PASSIVE.search(out) and not re.search(STATIVE, ol):
        add('L3 Semantic relations', 'voice (active -> passive)', PASSIVE.search(out).group(0))
    neg_usr = 'negative' in st or any(n['rel'] == 'neg' for n in nodes)
    if neg_usr and NEG_EN.search(ref) and not NEG_EN.search(out):
        add('L3 Semantic relations', 'negation lost', '')
    if not neg_usr and not NEG_EN.search(ref) and NEG_EN.search(out):
        add('L3 Semantic relations', 'negation added', NEG_EN.search(out).group(0))
    for n in nodes:
        pats = REL.get(n['rel'])
        if pats and has_any(ref, pats) and not has_any(out, pats):
            add('L3 Semantic relations', 'relation %s not expressed' % n['rel'], '/'.join(pats[:3]))
    # ---- L4 discourse
    for n in nodes:
        pats = DISC.get(n['disc'])
        if pats and has_any(ref, pats) and not has_any(out, pats):
            add('L4 Discourse', 'connective missing (%s)' % n['disc'], '/'.join(pats[:3]))
    # ---- L5 speaker's view
    seen = set()
    for n in nodes:
        for field in (n['spk'], n['scope']):
            for mk in re.findall(r'([A-Za-z]+)_\d+', field):
                pats = PART.get(mk)
                if pats and mk not in seen and has_any(ref, pats) and not has_any(out, pats):
                    seen.add(mk); add("L5 Speaker's view", 'particle %s lost' % mk, '/'.join(pats[:3]))
            for dx, pats in DEIX.items():
                if dx in (n['spk'], n['scope']) and has_any(ref, pats) and not has_any(out, pats) and dx not in seen:
                    other = DEIX['distal' if dx == 'proximal' else 'proximal']
                    seen.add(dx); add("L5 Speaker's view", 'deixis %s %s' % (dx, 'reversed' if has_any(out, other) else 'lost'), pats[0])
            if 'respect' in field and 'respect' not in seen and re.search(HONOR, rl) and not re.search(HONOR, ol):
                seen.add('respect'); add("L5 Speaker's view", 'honorific dropped', re.search(HONOR, rl).group(0))
    # ---- L6 sentence type
    q = 'interrogative' in st or hin.strip().endswith('?')
    if q and ref.strip().endswith('?') and not out.strip().endswith('?'):
        add('L6 Sentence type', 'question -> statement', '')
    if not q and out.strip().endswith('?') and not ref.strip().endswith('?'):
        add('L6 Sentence type', 'statement -> question', '')
    if 'imperative' in st and not re.search(r'\b(please|let|should|must)\b', ol) and re.match(r'^(you|it|this|the|there)\b', ol):
        add('L6 Sentence type', 'command -> statement', '')
    if hin.strip().endswith('!') and ref.strip().endswith('!') and not out.strip().endswith('!'):
        add('L6 Sentence type', 'exclamation lost', '')
    # ---- L7 adequacy (content coverage against the reference)
    rc = {stem(w) for w in toks(ref) if w not in STOP and len(w) > 2}
    oc = {stem(w) for w in toks(out) if w not in STOP and len(w) > 2}
    if rc:
        rec, prec = len(rc & oc) / len(rc), (len(rc & oc) / len(oc) if oc else 0)
        if rec < 0.35:
            add('L7 Adequacy', 'omission (low coverage)', 'recall %.2f; missing: %s' % (rec, ', '.join(sorted(rc - oc)[:6])))
        if prec < 0.35 and len(oc) >= 3:
            add('L7 Adequacy', 'addition (content not in reference)', 'precision %.2f; extra: %s' % (prec, ', '.join(sorted(oc - rc)[:6])))
    # ---- L8 output form
    if re.search(r'[ऀ-ॿ]', out):
        add('L8 Output form', 'Hindi script left', ' '.join(re.findall(r'[ऀ-ॿ]+', out))[:40])
    if re.search(r'\*\*|\(note|\bnote:|here is the translation|^(translation|english)\s*:', out, re.I):
        add('L8 Output form', 'notes / markdown', '')
    rep = re.search(r'\b(\w{3,}(?:\s+\w+){0,2})\s+\1\b', ol)
    if rep:
        add('L8 Output form', 'repetition', rep.group(0))
    lr, lo = len(ref.split()), len(out.split())
    if lr and lo > 2 * lr:
        add('L8 Output form', 'much longer than reference', '%d vs %d words' % (lo, lr))
    if lr >= 4 and lo < 0.5 * lr:
        add('L8 Output form', 'much shorter than reference', '%d vs %d words' % (lo, lr))
    return R


# ---------------- run ----------------
recs = []                                   # (tab, system, id, hindi, ref, out, subtype, detail)
full = collections.defaultdict(dict)
for it in items:
    i, ref = it['id'], it['reference']
    hin = hindi.get(i, it['hindi'])
    for s, o in SYS.items():
        out = o.get(i, '')
        R = analyse(it, hin, ref, out)
        full[s][i] = (hin, ref, out, R)
        for t, sub, det in R:
            recs.append((t, s, i, hin, ref, out, sub, det))

HDR = PatternFill('solid', fgColor='305496')
VERD = '"real error,acceptable variant,reference problem,not an error"'
wb = openpyxl.Workbook()
ov = wb.active; ov.title = 'Overview'
N = len(items)
ov.append(['%s: Experiment 1 layer-by-layer error analysis (USR-guided, reference-confirmed, no LLM judge)' % D.capitalize()])
ov['A1'].font = Font(bold=True, size=13); ov.append([])
ov.append(['Sentences per system', N]); ov.append([])
ov.append(['System', 'BLEU', 'chrF++', 'Sentences with any error'] + ['%s (sentences)' % t for t in TABS])
for s, o in SYS.items():
    hyp = [o.get(it['id'], '') for it in items]; ref = [it['reference'] for it in items]
    per = collections.defaultdict(set)
    for t, ss, i, *_ in recs:
        if ss == s:
            per[t].add(i)
    anyerr = set().union(*[per[t] for t in ERROR_TABS]) if per else set()
    ov.append([s, round(sacrebleu.corpus_bleu(hyp, [ref]).score, 2), round(sacrebleu.corpus_chrf(hyp, [ref], word_order=2).score, 2),
               '%d (%.1f%%)' % (len(anyerr), 100.0 * len(anyerr) / N)] + ['%d (%.1f%%)' % (len(per[t]), 100.0 * len(per[t]) / N) for t in TABS])
for c in ov[5]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = HDR
ov.append([])
ov.append(['Error sub-types (number of errors)'])
ov.cell(ov.max_row, 1).font = Font(bold=True)
ov.append(['Layer tab', 'Sub-type'] + list(SYS))
for c in ov[ov.max_row]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = HDR
cnt = collections.Counter((t, sub, s) for t, s, i, h, r, o, sub, det in recs)
for t in TABS:
    for sub in sorted({sub for tt, sub, s in cnt if tt == t}):
        ov.append([t, sub] + [cnt[(t, sub, s)] for s in SYS])
ov.append([])
for line in ['How to read: each layer has its own tab listing only that layer\'s errors (system, sentence, sub-type, detail).',
             'A flag is raised only when the USR and the gold reference agree and the output differs; confirm samples with the',
             '"Human verdict" column. L1c Synonyms lists acceptable variants (not errors). Cultural concepts are defined in',
             'resources/cultural_concepts.tsv and can be extended.']:
    ov.append([line])
ov.column_dimensions['A'].width = 26; ov.column_dimensions['B'].width = 30
for col in 'CDEFGHIJKLMNOP':
    ov.column_dimensions[col].width = 14

for t in TABS:
    ws = wb.create_sheet(t[:31])
    ws.append(['System', 'ID', 'Hindi', 'Gold reference', 'Output', 'Sub-type', 'Detail', 'Human verdict', 'Human note'])
    rows = sorted([r for r in recs if r[0] == t], key=lambda r: (r[6], r[1], r[2]))
    for r in rows:
        ws.append([r[1], r[2], r[3], r[4], r[5], r[6], r[7], '', ''])
    if rows:
        dv = DataValidation(type='list', formula1=VERD, allow_blank=True); ws.add_data_validation(dv); dv.add('H2:H%d' % (len(rows) + 1))
    for col, w in zip('ABCDEFGHI', [22, 9, 40, 40, 40, 26, 45, 16, 25]):
        ws.column_dimensions[col].width = w

al = wb.create_sheet('All sentences')
al.append(['ID', 'Hindi', 'Gold reference'] + [x for s in SYS for x in (s + ' output', s + ' errors')])
for it in items:
    i = it['id']
    row = [i, full[next(iter(SYS))][i][0], it['reference']]
    for s in SYS:
        hin, ref, out, R = full[s][i]
        row += [out, '; '.join('%s: %s %s' % (t.split()[0], sub, ('(' + det + ')') if det else '') for t, sub, det in R)]
    al.append(row)
for k, w in enumerate([9, 40, 40] + [40, 35] * len(SYS), 1):
    al.column_dimensions[openpyxl.utils.get_column_letter(k)].width = w

df = wb.create_sheet('Layer definitions')
for line in __doc__.strip().splitlines():
    df.append([line])
df.column_dimensions['A'].width = 130
for ws in wb.worksheets[1:]:
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = HDR
    if ws.title != 'Layer definitions':
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes = 'C2'
out = os.path.join(HERE, '%s_error_analysis.xlsx' % D)
wb.save(out)
print(out)
for t in TABS:
    print('%-24s' % t, ' '.join('%s=%d' % (s.split('_')[0], sum(1 for r in recs if r[0] == t and r[1] == s)) for s in SYS))
