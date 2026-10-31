"""J3: scores the independent third judgement (100 disputed candidates) and the author's check of the cause labels.
  third judgement: agreement with the adjudicated label, with the first and with the second judgement (raw, Cohen's
  kappa, Gwet's AC1), per layer; how often the third annotator sides with the adjudication when the two disagreed
  cause labels: share of the 100 keyword-assigned causes marked correct, and the corrections
Inputs: Desktop/annotator_tasks/14_J3_third_annotation/third_annotation.xlsx, cause_label_check.xlsx,
        third_annotation_KEY.tsv, cause_label_check_KEY.tsv, j3_final_labels.tsv
Output: j3_third_results.json"""
import collections, json, os, re, sys
import openpyxl
sys.stdout.reconfigure(encoding='utf-8')
HERE = os.path.dirname(os.path.abspath(__file__))
D = os.path.join(os.path.expanduser('~'), 'Desktop', 'annotator_tasks', '14_J3_third_annotation')
d1 = lambda s: int(re.match(r'\s*(\d)', str(s)).group(1)) if s is not None and re.match(r'\s*(\d)', str(s)) else None
KEY = {l.split('\t')[0]: l.rstrip('\n').split('\t') for l in list(open(os.path.join(HERE, 'third_annotation_KEY.tsv'), encoding='utf-8'))[1:]}
FL = {l.split('\t')[0]: l.rstrip('\n').split('\t') for l in list(open(os.path.join(HERE, 'j3_final_labels.tsv'), encoding='utf-8'))[1:]}
T = {}
for r in openpyxl.load_workbook(os.path.join(D, 'third_annotation.xlsx'))['Items'].iter_rows(min_row=2, values_only=True):
    if r[0]: T[r[0]] = d1(r[6])


def stats(pairs):
    """raw agreement, Cohen's kappa and Gwet's AC1 on binary labels (1 error, 2 not an error)"""
    n = len(pairs)
    if not n: return None
    po = sum(a == b for a, b in pairs) / n
    pa, pb = sum(a == 1 for a, _ in pairs) / n, sum(b == 1 for _, b in pairs) / n
    pe = pa * pb + (1 - pa) * (1 - pb)
    pi = (pa + pb) / 2; pe2 = 2 * pi * (1 - pi)
    return dict(n=n, raw=round(po, 3), kappa=round((po - pe) / (1 - pe), 3) if pe < 1 else None,
                ac1=round((po - pe2) / (1 - pe2), 3) if pe2 < 1 else None)


rows = []
for t, v in T.items():
    _, c, dno, layer = KEY[t]
    f = FL[c]
    rows.append(dict(item=t, cand=c, layer=f[1], third=v, first=d1(f[6]), second=d1(f[7]), final=d1(f[8])))
dec = [r for r in rows if r['third'] in (1, 2)]
res = dict(items=len(rows), third_verdicts=dict(collections.Counter(r['third'] for r in rows)),
           third_error_share=round(sum(r['third'] == 1 for r in dec) / len(dec), 3),
           final_error_share=round(sum(r['final'] == 1 for r in dec) / len(dec), 3))
res['vs_final'] = stats([(r['third'], r['final']) for r in dec if r['final'] in (1, 2)])
res['vs_first'] = stats([(r['third'], r['first']) for r in dec if r['first'] in (1, 2)])
res['vs_second'] = stats([(r['third'], r['second']) for r in dec if r['second'] in (1, 2)])
# when first and second gave opposite verdicts: with whom does the third side, and does it match the adjudication
opp = [r for r in dec if {r['first'], r['second']} == {1, 2}]
res['first_second_opposite'] = dict(n=len(opp), third_with_first=sum(r['third'] == r['first'] for r in opp),
                                    third_with_second=sum(r['third'] == r['second'] for r in opp),
                                    third_with_final=sum(r['third'] == r['final'] for r in opp))
res['by_layer'] = {}
for L in sorted({r['layer'] for r in rows}):
    sub = [r for r in dec if r['layer'] == L]
    res['by_layer'][L] = dict(n=len(sub), third_error=sum(r['third'] == 1 for r in sub), final_error=sum(r['final'] == 1 for r in sub),
                              agree_final=sum(r['third'] == r['final'] for r in sub))
# cause labels -----------------------------------------------------------------------------------------------------
CK = {l.split('\t')[0]: l.rstrip('\n').split('\t') for l in list(open(os.path.join(HERE, 'cause_label_check_KEY.tsv'), encoding='utf-8'))[1:]}
ok, wrong, blank = 0, [], 0
for r in openpyxl.load_workbook(os.path.join(D, 'cause_label_check.xlsx'))['Reasons'].iter_rows(min_row=2, values_only=True):
    if not r[0]: continue
    v = str(r[7] or '').strip().lower()
    if v.startswith('y'): ok += 1
    elif v.startswith('n'): wrong.append((r[0], r[6], r[8]))
    else: blank += 1
res['cause_labels'] = dict(checked=ok + len(wrong), correct=ok, share_correct=round(ok / max(1, ok + len(wrong)), 3),
                           blank=blank, corrections=collections.Counter('%s -> %s' % (a, b) for _, a, b in wrong).most_common())
json.dump(res, open(os.path.join(HERE, 'j3_third_results.json'), 'w', encoding='utf-8'), indent=1, ensure_ascii=False)
print(json.dumps(res, indent=1, ensure_ascii=False))
