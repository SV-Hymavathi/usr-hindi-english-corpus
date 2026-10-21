#!/usr/bin/env python3
"""Corpus statistics of the gold USR dataset (tables of the LRE paper).
    python dataset_stats.py            -> dataset_stats.xlsx, dataset_stats.tsv
Reads ../../corpus/<domain>.jsonl (gold USR, Hindi, reference) and ../../usr2chunk/data/concept-to-mrs-rels.dat."""
import collections, json, os, re
import openpyxl
from openpyxl.styles import Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
REL = os.path.join(HERE, '..', '..')
DOMS = ['tourism', 'health', 'geography']
REL_GROUP = {'k1': 'karta (agent/subject)', 'k2': 'karma (object)', 'k3': 'karana (instrument)', 'k4': 'sampradana (recipient)',
             'k5': 'apadana (source)', 'k7p': 'location (place)', 'k7t': 'location (time)', 'k7': 'location (other)',
             'r6': 'genitive', 'mod': 'modifier', 'k1s': 'subject complement', 'k2s': 'object complement', 'rh': 'cause',
             'rt': 'purpose', 'card': 'cardinal', 'ord': 'ordinal', 'quant': 'quantifier', 'dem': 'demonstrative',
             'intf': 'intensifier', 'neg': 'negation', 'krvn': 'manner adverbial', 'rvks': 'participle (V-ing) modifier',
             'rbks': 'participle (V-ed) modifier', 'rpk': 'conjunctive participle', 'rsm': 'similarity', 'rasneg': 'without'}
stat, rel, tam, disc, cx, stype, spk = {}, {}, {}, {}, {}, {}, {}
for d in DOMS:
    items = [json.loads(l) for l in open(os.path.join(REL, 'corpus', d + '.jsonl'), encoding='utf-8')]
    rel[d], tam[d], disc[d], cx[d], stype[d], spk[d] = (collections.Counter() for _ in range(6))
    n_nodes = n_concepts = n_names = 0
    base = set(); toks = 0; ref_toks = 0; src_len = collections.Counter(); usr_len = []
    for it in items:
        base.add(re.sub(r'(?<=\d)[a-z]$|(?<=\d)[a-z]\d?$', '', it['id']))
        toks += len(it['hindi'].split()); ref_toks += len(it['reference'].split())
        words = len([w for w in it['hindi'].split() if re.search(r'\w', w)])   # punctuation tokens not counted
        usr_len.append(words); src_len[re.sub(r'(?<=\d)[a-z]$|(?<=\d)[a-z]\d?$', '', it['id'])] += words
        for l in it['usr'].splitlines():
            l = l.strip()
            if l.startswith('%'):
                stype[d][l] += 1; continue
            c = l.split()
            if len(c) < 9:
                continue
            n_nodes += 1
            if c[0].startswith('['):
                cx[d][re.sub(r'_\d+\]$', ']', c[0])] += 1
            elif re.search(r'_\d+', c[0]):
                n_concepts += 1
            if ':' in c[4]:
                rel[d][c[4].split(':')[1]] += 1
            m = re.match(r'^.*?_\d+-(.+)$', c[0])
            if m:
                tam[d][m.group(1)] += 1
            if ':' in c[5]:
                disc[d][c[5].split(':')[1]] += 1
            for f in (c[6], c[7]):
                for mk in re.findall(r'[A-Za-z]+(?:_\d+)?', f):
                    spk[d][mk] += 1
            if re.search(r':(begin|inside)$', c[8]):
                n_names += 1
    stat[d] = dict(usr=len(items), source=len(base), nodes=n_nodes, concepts=n_concepts, name_tokens=n_names,
                   hindi_tokens=toks, ref_tokens=ref_toks, avg_nodes=round(n_nodes / len(items), 2),
                   passive=sum(v for k, v in stype[d].items() if 'pass' in k),
                   words_per_source_sentence=round(sum(src_len.values()) / len(src_len), 1),
                   max_words_source_sentence=max(src_len.values()),
                   words_per_usr=round(sum(usr_len) / len(usr_len), 1), max_words_usr=max(usr_len))
dict_new = collections.Counter()
for l in open(os.path.join(REL, 'usr2chunk', 'data', 'concept-to-mrs-rels.dat'), encoding='utf-8'):
    c = l.split()
    if c and re.search(r'_9[0-9]$', c[0]):
        dict_new['complex predicate (cp)' if '+' in c[0] else 'concept sense'] += 1

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Overview'
HDR = PatternFill('solid', fgColor='305496')
def hdr(w):
    for c in w[w.max_row]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = HDR
ws.append(['Statistic'] + DOMS + ['Total']); hdr(ws)
for k, lab in [('source', 'Source sentences'), ('usr', 'USRs (after clause segmentation)'), ('hindi_tokens', 'Hindi tokens'),
               ('ref_tokens', 'English reference tokens'), ('nodes', 'USR rows (nodes)'), ('concepts', 'Concepts with dictionary sense'),
               ('name_tokens', 'Named-entity tokens'), ('avg_nodes', 'Mean nodes per USR'), ('passive', 'Passive USRs'),
               ('words_per_source_sentence', 'Mean Hindi words per source sentence'), ('words_per_usr', 'Mean Hindi words per USR')]:
    vals = [stat[d][k] for d in DOMS]
    ws.append([lab] + vals + ['' if k.startswith(('avg', 'words')) else sum(vals)])
for title, C, top in [('Relations', rel, 30), ('TAM', tam, 30), ('Discourse', disc, 20), ('Constructions', cx, 15),
                      ('Sentence types', stype, 15), ("Speaker's view", spk, 20)]:
    w = wb.create_sheet(title); w.append([title[:-1] if title.endswith('s') else title] + DOMS + ['Total'] +
                                        (['Description'] if title == 'Relations' else [])); hdr(w)
    tot = collections.Counter()
    for d in DOMS:
        tot.update(C[d])
    for k, v in tot.most_common(top):
        w.append([k] + [C[d][k] for d in DOMS] + [v] + ([REL_GROUP.get(k, '')] if title == 'Relations' else []))
    w.column_dimensions['A'].width = 28
dw = wb.create_sheet('Dictionary additions'); dw.append(['Type', 'New entries (sense number 90-99)']); hdr(dw)
for k, v in dict_new.items():
    dw.append([k, v])
ws.column_dimensions['A'].width = 34
wb.save(os.path.join(HERE, 'dataset_stats.xlsx'))
with open(os.path.join(HERE, 'dataset_stats.tsv'), 'w', encoding='utf-8') as f:
    for d in DOMS:
        f.write(d + '\t' + json.dumps(stat[d]) + '\n')
    f.write('dictionary_additions\t' + json.dumps(dict_new) + '\n')
for d in DOMS:
    print(d, stat[d])
print('top relations', (rel['tourism'] + rel['health'] + rel['geography']).most_common(10))
print('discourse', (disc['tourism'] + disc['health'] + disc['geography']).most_common(8))
print('constructions', (cx['tourism'] + cx['health'] + cx['geography']).most_common(8))
print('dict', dict_new)
