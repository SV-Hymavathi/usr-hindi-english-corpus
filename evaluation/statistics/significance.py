#!/usr/bin/env python3
"""Paired bootstrap significance tests (Koehn 2004) for BLEU and chrF++ between systems.

Reproduce:
    pip install sacrebleu==2.6.0 numpy openpyxl
    python significance.py                 # -> significance_results.xlsx and significance_results.tsv
    python significance.py --samples 1000 --seed 12345

Data read (relative to this folder):
    ../inputs/<domain>.jsonl              gold reference translation for every USR id (field 'reference')
    ../exp1/outputs/<system>.tsv          Experiment 1 (Hindi -> English): domain <TAB> id <TAB> output
    ../exp3/outputs/<system>.tsv          Experiment 3 (generator chunks -> fluent sentence)
    ../exp2/outputs/<system>.tsv          Experiment 2 (USR -> sentence), included automatically when complete

Method: for every pair (A, B) on the same sentences, draw N bootstrap samples of sentence indices (with replacement,
same indices for A and B), compute corpus BLEU / chrF++ of A and B on each sample from per-sentence sufficient statistics
(sacrebleu), and report p = share of samples in which the system with the lower full-corpus score is at least as good as
the other (two systems differ significantly at p < 0.05). 95% confidence intervals of each system's score are the
2.5 / 97.5 percentiles of its bootstrap scores. Metric settings: sacrebleu defaults (BLEU: 13a tokenizer, exp smoothing;
chrF++: char order 6, word order 2, beta 2)."""
import argparse, itertools, json, os
import numpy as np
from sacrebleu.metrics import BLEU, CHRF
import openpyxl
from openpyxl.styles import Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
EXP = os.path.join(HERE, '..')
ap = argparse.ArgumentParser()
ap.add_argument('--samples', type=int, default=1000)
ap.add_argument('--seed', type=int, default=12345)
a = ap.parse_args()
DOMS = ['tourism', 'health', 'geography']
REF = {}
for d in DOMS:
    for l in open(os.path.join(EXP, 'inputs', d + '.jsonl'), encoding='utf-8'):
        r = json.loads(l)
        REF[(d, r['id'])] = r['reference']
KEYS = [k for k in REF if REF[k]]


def load(exp, name):
    f = os.path.join(EXP, exp, 'outputs', name + '.tsv')
    if not os.path.exists(f):
        return None
    o = {}
    for l in open(f, encoding='utf-8'):
        p = l.rstrip('\n').split('\t', 2)
        if len(p) == 3:
            o[(p[0], p[1])] = p[2]
    return o if all(k in o for k in KEYS) else None          # only complete systems


SYSTEMS = {}
for exp, names in (('exp1', ['google_translate', 'indictrans2_local', 'groq_gpt-oss-120b', 'mistral_ministral-14b', 'ollama_gemma3-4b']),
                   ('exp3', ['groq_gpt-oss-120b', 'mistral_ministral-14b', 'ollama_gemma3-4b', 'usr2chunk_literal']),
                   ('exp2', [m + c for m in ('groq_gpt-oss-120b', 'mistral_ministral-14b', 'ollama_gemma3-4b')
                             for c in ('_raw', '_guide', '_gloss')])):
    for n in names:
        o = load(exp, n)
        if o:
            SYSTEMS['%s:%s' % (exp, n)] = o
print('systems:', list(SYSTEMS))

METRICS = {'BLEU': BLEU(), 'chrF++': CHRF(word_order=2)}
rng = np.random.default_rng(a.seed)
results, ci_rows = [], []
for scope in ['all'] + DOMS:
    keys = [k for k in KEYS if scope == 'all' or k[0] == scope]
    refs = [[REF[k] for k in keys]]
    idx = rng.integers(0, len(keys), size=(a.samples, len(keys)))
    stats = {}
    for mname, m in METRICS.items():
        for s, o in SYSTEMS.items():
            st = np.array(m._extract_corpus_statistics([o[k] for k in keys], refs), dtype=float)
            full = m._compute_score_from_stats(list(st.sum(0))).score
            boot = np.array([m._compute_score_from_stats(list(st[ix].sum(0))).score for ix in idx])
            stats[(mname, s)] = (full, boot)
            ci_rows.append([scope, mname, s, round(full, 2), round(float(np.percentile(boot, 2.5)), 2), round(float(np.percentile(boot, 97.5)), 2)])
    for mname in METRICS:
        for A, B in itertools.combinations(SYSTEMS, 2):
            fa, ba = stats[(mname, A)]; fb, bb = stats[(mname, B)]
            better, worse = (A, B) if fa >= fb else (B, A)
            bw = stats[(mname, worse)][1]; bb_ = stats[(mname, better)][1]
            p = float(np.mean(bw >= bb_))
            results.append([scope, mname, A, round(fa, 2), B, round(fb, 2), round(fa - fb, 2), better, round(p, 4),
                             'yes' if p < 0.05 else 'no'])
    print(scope, 'done')

with open(os.path.join(HERE, 'significance_results.tsv'), 'w', encoding='utf-8') as f:
    f.write('scope\tmetric\tsystem_A\tscore_A\tsystem_B\tscore_B\tdiff\tbetter\tp_value\tsignificant_p<0.05\n')
    for r in results:
        f.write('\t'.join(map(str, r)) + '\n')
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Pairwise tests'
HDR = PatternFill('solid', fgColor='305496')
ws.append(['Scope', 'Metric', 'System A', 'Score A', 'System B', 'Score B', 'A - B', 'Better', 'p-value', 'Significant (p<0.05)'])
for r in results:
    ws.append(r)
c2 = wb.create_sheet('Scores with 95% CI')
c2.append(['Scope', 'Metric', 'System', 'Score', 'CI low', 'CI high'])
for r in ci_rows:
    c2.append(r)
info = wb.create_sheet('Method')
for line in __doc__.strip().splitlines() + ['', 'samples = %d, seed = %d' % (a.samples, a.seed)]:
    info.append([line])
for w in (ws, c2):
    for c in w[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = HDR
    for col, wd in zip('ABCDEFGHIJ', [10, 8, 36, 9, 36, 9, 8, 36, 9, 12]):
        w.column_dimensions[col].width = wd
info.column_dimensions['A'].width = 130
wb.save(os.path.join(HERE, 'significance_results.xlsx'))
print('written: significance_results.xlsx / .tsv  (%d tests)' % len(results))
